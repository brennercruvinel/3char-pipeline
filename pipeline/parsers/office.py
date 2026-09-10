"""xlsx workbooks and the fixed width ICD order file."""
import pathlib, re, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from core import parser, row, ROOT

@parser("cls/icd")
def icd10cm(path, sysrow):
    """CMS/NCHS order file, fixed width: order number, code, a billable flag, a 60
    character short description, then the long one. flag 0 means the code is a heading
    that cannot be billed on its own, which is the difference between a category and a
    diagnosis you can actually put on a claim."""
    text = (ROOT/path).read_text(encoding="latin-1")
    for line in text.splitlines():
        if len(line) < 20: continue
        code = line[6:13].strip()
        if not code: continue
        billable = line[14:15].strip() == "1"
        short = line[16:76].strip()
        long_ = line[77:].strip()
        parent = code[:-1] if len(code) > 3 else (code[:3] if len(code) > 3 else None)
        yield row("cls/icd", code, long_ or short,
                  description=short if short != long_ else None,
                  parent_code=parent if parent and parent != code else None,
                  extra={"billable": billable, "order": line[:5].strip(),
                         "short_description": short})

def _sheet_rows(wb, name):
    ws = wb[name]
    return [list(r) for r in ws.iter_rows(values_only=True)]

@parser("fin/cfi")
def cfi(path, sysrow):
    """ISO 10962 is positional: six characters, category, group, then four attributes.
    the workbook writes every level as a masked code in the rdfs label column, EXXXXX
    for the category, ESXXXX for the group, ESVXXX once the first attribute is set.
    the mask is the hierarchy, so the parent is this code with its last real character
    turned back into X."""
    import openpyxl, warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(ROOT/path, read_only=True, data_only=True)
    seen = {}
    order = []
    for name in wb.sheetnames:
        if name == "Introduction": continue
        for r in _sheet_rows(wb, name):
            cells = [str(c).strip() if c is not None else "" for c in r]
            label = next((c for c in cells if re.fullmatch(r"[A-Z]{6}", c) and "X" in c), None)
            if not label or label in seen: continue
            texts = [c for c in cells if c and c != label and not re.fullmatch(r"[A-Z\-\*]{6,11}", c)]
            nm = texts[0] if texts else ""
            desc = texts[1] if len(texts) > 1 else None
            order.append(label)
            seen[label] = (nm, desc)
    wb.close()
    for label in order:
        nm, desc = seen[label]
        real = label.rstrip("X")
        parent = (real[:-1] + "X" * (6 - len(real) + 1)) if len(real) > 1 else None
        yield row("fin/cfi", label, nm, description=desc,
                  parent_code=parent if parent in seen else None,
                  extra={"depth": len(real), "mask": label})

@parser("fin/mic")
def mic(path, sysrow):
    """ISO 10383. an operating MIC runs a market, a segment MIC is one book inside it,
    and OPRT/SGMT says which this row is."""
    import openpyxl, warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(ROOT/path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(head)}
    g = lambda r, k: (str(r[idx[k]]).strip() if k in idx and idx[k] < len(r)
                      and r[idx[k]] is not None else None)
    for r in rows:
        code = g(r, "MIC")
        if not code: continue
        oprt = g(r, "OPERATING MIC")
        kind = g(r, "OPRT/SGMT")
        status = (g(r, "STATUS") or "").lower() or None
        yield row("fin/mic", code,
                  g(r, "MARKET NAME-INSTITUTION DESCRIPTION") or g(r, "LEGAL ENTITY NAME") or "",
                  status="deprecated" if status and status != "active" else None,
                  parent_code=oprt if kind == "SGMT" and oprt != code else None,
                  aliases=[a for a in (g(r, "ACRONYM"),) if a],
                  extra={"kind": kind, "operating_mic": oprt, "lei": g(r, "LEI"),
                         "market_category": g(r, "MARKET CATEGORY CODE"),
                         "country": g(r, "ISO COUNTRY CODE (ISO 3166)"),
                         "city": g(r, "CITY"), "status": status})
    wb.close()
